from pptx import Presentation
from pptx.util import Inches, Pt

def create_respiratory_shield_pptx():
    prs = Presentation()

    # Helper function to add content slides
    def add_content_slide(title_text, points):
        slide_layout = prs.slide_layouts[1] # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = title_text
        
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        for point in points:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0
        return slide

    # --- Slide 1: Title Slide (Replicating Template Structure) ---
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "DEVELOPMENT AND IMPLEMENTATION OF A LOW-POWER MEMS-BASED WEARABLE RESPIRATORY SHIELD FOR REAL-TIME ASTHMA TRIGGER DETECTION AND EXPOSURE MAPPING"
    subtitle.text = ("FEDERAL UNIVERSITY OF TECHNOLOGY, MINNA\n"
                     "Department of Computer Engineering\n"
                     "2024/2025")

    # --- Slide 2: Outline ---
    add_content_slide("Outline", [
        "1. Introduction & Background",
        "2. Problem Statement",
        "3. Research Objectives",
        "4. Literature Review",
        "5. Significance of the Study",
        "6. Scope & Limitations"
    ])

    # --- Slide 3: Introduction & Background ---
    add_content_slide("Introduction & Background", [
        "Asthma affects ~262 million people globally, with a disproportionate burden in sub-Saharan Africa[cite: 383, 385].",
        "Traditional management relies on reactive, post-symptom intervention[cite: 387].",
        "MEMS technology enables miniaturized, low-power sensing of environmental triggers[cite: 390, 391].",
        "The project integrates MEMS sensing, IoT, and health informatics for real-time alerts."
    ])

    # --- Slide 4: Problem Statement ---
    add_content_slide("Problem Statement", [
        "High rates of uncontrolled asthma attacks in Nigeria due to a lack of personal monitoring tools[cite: 397, 398].",
        "Commercial solutions are often financially and geographically inaccessible[cite: 399].",
        "Government environmental monitoring is too sparse to capture individual micro-exposures[cite: 400].",
        "Critical gap between MEMS technical potential and practical availability for local patients[cite: 401]."
    ])

    # --- Slide 5: Research Objectives ---
    add_content_slide("Research Objectives", [
        "Aim: Develop a low-power wearable shield for real-time trigger detection and exposure mapping[cite: 403].",
        "Characterize primary asthma triggers relevant to the Niger State context[cite: 404].",
        "Integrate MEMS sensors for PM2.5/10, VOCs, humidity, and temperature[cite: 405].",
        "Design a low-power embedded system for continuous real-time data acquisition[cite: 406].",
        "Implement GPS-referenced data logging for exposure mapping[cite: 408]."
    ])

    # --- Slide 6: Literature Review Summary ---
    add_content_slide("Literature Review (Chapter 2)", [
        "Review of asthma epidemiology from international and Nigerian perspectives[cite: 426].",
        "Analysis of MEMS sensor technology and wearable health monitoring systems[cite: 426].",
        "Identified gap: Limited locally developed research on low-power embedded respiratory monitors[cite: 402].",
        "Focus on integrating low-power wireless protocols and edge computing[cite: 392]."
    ])

    # --- Slide 7: Significance of the Study ---
    add_content_slide("Significance of the Study", [
        "Engineering: Demonstrates local hardware prototyping using commercial MEMS components[cite: 415].",
        "Clinical: Addresses unmet needs for personalized monitoring in Nigerian communities[cite: 416].",
        "Public Health: Generates granular, geo-referenced air quality data for epidemiological research[cite: 417].",
        "Academic: Provides a reference design for future IoT-based health systems at FUT-Minna[cite: 418]."
    ])

    # --- Slide 8: Scope and Limitations ---
    add_content_slide("Scope & Limitations", [
        "Scope: Prototype development and laboratory/field evaluation in Minna metropolis[cite: 419, 422].",
        "Targets: PM, VOCs, humidity, and temperature as primary trigger proxies[cite: 420].",
        "Limitation: Clinical validation with confirmed patients is reserved for future work[cite: 421].",
        "Note: System complements, rather than replaces, medical diagnosis[cite: 423]."
    ])

    file_name = "Respiratory_Shield_Project_Presentation.pptx"
    prs.save(file_name)
    print(f"Presentation saved as {file_name}")

if __name__ == "__main__":
    create_respiratory_shield_pptx()